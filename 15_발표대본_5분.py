# -*- coding: utf-8 -*-
"""
15. 5분 발표 대본(워드, 가로)을 만든다 — 최종 발표자료(output/한눈_발표자료.pdf, 19장) 기준

- 한계 슬라이드는 발표자료에서 빠졌으므로 대본에도 없다. 19쪽 부록은 질문이 나올 때만.
- 표: 쪽 | 제목 | 시간 | 말할 내용 | 동작(가리킬 곳)
- 분량: 한국어 발표 속도 약 분당 330~350자(공백 제외) 기준으로 5분 안팎

산출: output/한눈_발표대본_5분.docx
"""

import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "output", "한눈_발표대본_5분.docx")

# (쪽, 제목, 초, 말할 내용, 동작)
ROWS = [
    (1, "표지", 10,
     "안녕하십니까. 휴먼AI공학전공 김성빈입니다. 번역 없이 문화를 비교하는 AI 비전 팀플 도구, '한눈'을 소개하겠습니다.",
     "인사"),
    (2, "목차 · 발표자", 10,
     "문제 인식, 해결 방안, 파일럿 실험, 프로토타입, 기대효과 순서로 말씀드리겠습니다.",
     "목차 훑기"),
    (3, "문제 인식", 25,
     "유학생과 팀플을 해 보셨다면 익숙한 장면일 겁니다. 한국어가 서툰 팀원에게 무엇을 맡길지 애매하고, 그 팀원이 모국어로 "
     "찾아온 자료는 한국 자료와 이어지지 않습니다. 그래서 유학생의 조사는 결국 '해외 사례' 한 장으로 떨어져 나갑니다.",
     "왼쪽 목차의 '4. 해외 사례' 가리키기"),
    (4, "현황", 25,
     "국내 대학 유학생은 30만 7천 명으로 3년 새 1.7배가 됐지만, 어학 기준을 충족한 학생은 49.9%, 절반뿐입니다. 유학생의 "
     "36.6%는 팀 프로젝트 참여가 어렵다고 답했습니다. 같은 강의실에 있지만 섞이지 못하는 두 개의 섬입니다.",
     "숫자 세 개를 왼쪽부터"),
    (5, "문제 정의", 20,
     "번역을 거친 자료는 타국 이야기로 읽히고, 한국 자료와 비교 기준이 다르며, 발표 직전에 이어 붙이니 논리가 이어지지 않습니다. "
     "핵심은 언어 실력이 아니라, 서로 다른 언어의 자료가 같은 기준 위에 놓이지 못하는 구조입니다.",
     "아래 파란 상자 강조"),
    (6, "번역이 아닌 시각적 해석", 15,
     "그래서 발상을 바꿨습니다. 사진과 영상은 번역이 필요 없습니다. 측정은 AI가 하고, 해석은 두 문화의 사람이 합니다.",
     "'측정'·'해석' 카드"),
    (7, "해결 방안", 25,
     "한눈은 네 단계입니다. 팀원이 자기 나라 현장 사진을 직접 찾아 올리면, YOLO가 사람과 의자, 오토바이를 찾아 세고 생활 지표로 "
     "바꿉니다. 나라별 지표를 같은 차트에 놓고, 차이가 큰 부분은 그 나라 출신 팀원이 이유를 설명합니다.",
     "①→④ 순서대로"),
    (8, "AI 설계", 15,
     "탐지 결과를 좌석, 착석, 무리, 이동 수단이라는 네 가지 지표로 바꾸는 규칙은 직접 설계했습니다. 사람 박스 아래쪽이 의자와 "
     "겹치면 앉아 있다고 보는 식입니다.",
     "사진 속 '앉은 사람' 박스"),
    (9, "파일럿 실험", 20,
     "실제로 해 봤습니다. 팀원이 올릴 자료 대신 한국, 일본, 베트남 편의점 영상 21편을 사례로 쓰고, 796개 장면을 같은 기준으로 "
     "쟀습니다. 가설은 한국은 사고 나가는 곳, 베트남은 머무는 곳이었습니다.",
     "'규모'·'가설' 줄"),
    (10, "검증", 10,
     "자료를 고르고 걸러내는 일은 사람이 직접 했습니다. AI는 걸러진 자료를 잴 뿐입니다.",
     "제외한 영상 목록"),
    (11, "결과 ①", 15,
     "좌석, 앉은 사람, 무리, 오토바이 모두 베트남에서 가장 자주 보였습니다. 일본은 오히려 한국보다 낮았습니다.",
     "레이더 차트 → 베트남 카드"),
    (12, "결과 ②", 20,
     "다만 영상별로 보면 베트남의 높은 값은 7편 중 2편에서 나왔습니다. 그래서 한눈은 결론 대신 질문을 줍니다. 이 매장들은 왜 "
     "달랐을까. 그 답은 현지 팀원이 할 수 있습니다.",
     "베트남 막대 위 높은 점 두 개"),
    (13, "검수", 15,
     "AI가 맞게 봤는지 한 장씩 채점했습니다. 사람은 83%, 좌석은 78%를 맞혔고, 틀린 좌석은 고쳐 반영해도 순위는 같았습니다.",
     "83%·78%"),
    (14, "결과 ③", 10,
     "실제로 AI가 본 장면입니다. 파란 박스가 좌석, 주황 박스가 사람입니다.",
     "베트남 사진"),
    (15, "해석 카드", 20,
     "AI는 무엇이 얼마나 보였는지만 말합니다. 왜 그런지는 현지 팀원이 해석하고, 그 해석이 우리 과제의 주장이 됩니다. 외국 "
     "자료가 부록이 아니라 결론의 근거가 되는 순간입니다.",
     "세 카드 왼쪽→오른쪽"),
    (16, "프로토타입", 15,
     "이 흐름을 5단계 웹 프로토타입으로 직접 만들었습니다. 측정, 비교, 해석 카드, 그리고 발표 본론으로 보내기까지 이어집니다.",
     "네 화면 순서대로 (시연 영상은 별도 제출)"),
    (17, "기대효과", 20,
     "문화 차이는 번역하며 사라지는 잡음이 아니라 해석할 인사이트가 되고, 부록이던 외국 자료는 본론이 됩니다. 편의점은 예시일 "
     "뿐, 지표 규칙만 바꾸면 거리나 공원, 캠퍼스 등 어떤 현장 비교에도 쓸 수 있습니다.",
     "아래 '확장성' 상자"),
    (18, "마무리", 10,
     "팀플의 해외 사례 한 장이 본론이 되는 순간, 같은 강의실의 두 섬은 하나의 팀이 됩니다. 감사합니다.",
     "인사"),
    (19, "부록", 0,
     "(질문이 나올 때만) 분석에 쓴 공개 영상 목록입니다.",
     "질문 시"),
]


def shade(cell, hex_):
    tc = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_)
    tc.append(s)


def run(p, text, size=11, bold=False, color=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.4)

    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(11)

    total = sum(r[2] for r in ROWS)
    chars = sum(len(r[3].replace(" ", "")) for r in ROWS if r[2])
    p = doc.add_paragraph()
    run(p, "「한눈」 발표 대본 (5분)", 18, True)
    p = doc.add_paragraph()
    run(p, "2026 HUSS · AI를 활용한 어울림 공모전  |  발표 김성빈  |  슬라이드 19장  |  예상 {}분 {}초 (공백 제외 {}자)".format(
        total // 60, total % 60, chars), 10, color="6B7684")

    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.2), Cm(3.4), Cm(1.6), Cm(15.3), Cm(5.0)]
    for i, h in enumerate(["쪽", "슬라이드", "시간", "말할 내용", "동작 · 가리킬 곳"]):
        c = t.rows[0].cells[i]
        c.width = widths[i]
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(c.paragraphs[0], h, 10, True, "FFFFFF")
        shade(c, "3182F6")
    # 머리행 반복
    tr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    tr.append(th)

    acc = 0
    for no, title, sec_, text, act in ROWS:
        cells = t.add_row().cells
        for i, w in enumerate(widths):
            cells[i].width = w
        acc += sec_
        vals = [str(no), title, "{}초".format(sec_) if sec_ else "—", text, act]
        for i, v in enumerate(vals):
            para = cells[i].paragraphs[0]
            if i in (0, 2):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(para, v, 12 if i == 3 else 10, bold=(i == 1), color=("8B95A1" if i == 4 else None))
            para.paragraph_format.line_spacing = 1.35 if i == 3 else 1.1
        if sec_:
            q = cells[2].add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(q, "~{}:{:02d}".format(acc // 60, acc % 60), 8, color="8B95A1")
        if no % 2 == 0:
            for c in cells:
                shade(c, "F2F4F6")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run(p, "팁: 시간이 모자라면 10쪽(검증)·14쪽(결과 ③)은 한 문장만 말하고 넘어가도 흐름이 이어집니다. "
           "'~0:00'은 그 쪽을 마칠 때까지의 누적 시간입니다.", 9, color="6B7684")
    doc.save(OUT)
    print("[저장]", OUT, "{}초, {}자".format(total, chars))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
