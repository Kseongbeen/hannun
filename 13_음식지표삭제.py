# -*- coding: utf-8 -*-
"""
13. 음식·음료 인식(검수 정밀도 0%)을 문서 전체에서 뺀다 — 사용자가 세 번째로 손본 기획서 기준

- 기획서(원본: output/한눈_기획서_사용자수정3_백업.hwpx)
  · 2-3 (3) '테이블 위 음식' 제외 문장, (5) 시사점의 '음식처럼…' 문장, [표 8]의 '음식·음료' 행 삭제
  · YOLO 각주의 예시 '사람·의자·컵' → '사람·의자·오토바이'
  · 사용자가 1-1 '고립' 항목(오마이뉴스)을 지웠으므로 참고 자료에서도 오마이뉴스 삭제
  · [그림 1] 프로토타입 화면을 음식 범례가 빠진 새 캡처로 교체
- 발표(원본: output/한눈_발표.pptx): 13쪽 '음식·음료' 행 도형 삭제, 18쪽 '음식은 0%라 제외' 문구 삭제,
  16쪽 프로토타입 캡처 교체, 13쪽 노트(대본) 수정
- 대본: 13쪽 대본에서 음식 문장 삭제 → docx 다시 만듦
- 덮어쓰기 전 기존 파일은 output/_자동백업/에 남긴다.
"""

import importlib.util
import json
import os
import re
import sys
import zipfile
import xml.dom.minidom

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "output")
_spec = importlib.util.spec_from_file_location("m12", os.path.join(HERE, "12_개념반영_대본.py"))
m12 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(m12)

HWP_SRC = os.path.join(OUT, "한눈_기획서_사용자수정3_백업.hwpx")
HWP_DST = os.path.join(OUT, "한눈_기획서.hwpx")
PPT = os.path.join(OUT, "한눈_발표.pptx")
MATCH = json.load(open(os.path.join(OUT, "_render", "img_match.json"), encoding="utf-8"))   # 문서 속 그림 ↔ 캡처 파일

HWP_EDITS = [
    (". '테이블 위 음식'도 검수 정밀도 0%로 제외((5) 참고). 좌석 오답", ". 좌석 오답"),
    ("보정으로 나라 간 차이가 오히려 조금 커짐. 음식처럼 정밀도가 낮은 지표는 결과에서 빼는 것이 원칙",
     "보정으로 나라 간 차이가 오히려 조금 커짐"),
    ("사람·의자·컵 등 80종", "사람·의자·오토바이 등 80종"),
]


def cut(x, start_tag, end_tag, needle):
    k = x.find(needle)
    assert k > 0, needle
    a = x.rfind(start_tag, 0, k)
    b = x.find(end_tag, k) + len(end_tag)
    return x[:a] + x[b:], x[a:b]


def build_hwp():
    z = zipfile.ZipFile(HWP_SRC)
    x = z.read("Contents/section0.xml").decode("utf-8")
    for a, b in HWP_EDITS:
        assert x.count(a) == 1, a
        x = x.replace(a, b)
    # [표 8] 마지막 행(음식·음료) 삭제
    k = x.find("<hp:t>음식·음료</hp:t>")
    t0 = x.rfind("<hp:tbl ", 0, k)
    x, row = cut(x, "<hp:tr>", "</hp:tr>", "<hp:t>음식·음료</hp:t>")
    h = max(int(v) for v in re.findall(r'<hp:cellSz width="\d+" height="(\d+)"', row))
    head = x[t0:x.find(">", t0) + 1]
    n = int(re.search(r'rowCnt="(\d+)"', head).group(1))
    x = x[:t0] + head.replace('rowCnt="{}"'.format(n), 'rowCnt="{}"'.format(n - 1)) + x[t0 + len(head):]
    m = re.compile(r'(<hp:sz width="\d+" widthRelTo="ABSOLUTE" height=")(\d+)"').search(x, t0)
    x = x[:m.start(2)] + str(int(m.group(2)) - h) + x[m.end(2):]
    # 참고 자료의 오마이뉴스(본문 '고립' 항목을 사용자가 지움)
    x, _ = cut(x, "<hp:p ", "</hp:p>", "· 오마이뉴스(2020.12.1.)")
    assert "음식" not in re.sub(r"<[^>]+>", "", x).replace("음식 클로즈업", ""), "음식 문구가 남음"
    x = re.sub(r"<hp:linesegarray>.*?</hp:linesegarray>", "", x, flags=re.S)
    xml.dom.minidom.parseString(x.encode("utf-8"))
    new_img = {n: open(os.path.join(OUT, f), "rb").read() for n, f in MATCH["hwpx"].items()}
    tmp = HWP_DST + ".tmp"
    with zipfile.ZipFile(tmp, "w") as out:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename == "Contents/section0.xml":
                data = x.encode("utf-8")
            elif info.filename in new_img:
                data = new_img[info.filename]
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = info.compress_type
            out.writestr(zi, data)
    z.close()
    os.replace(tmp, HWP_DST)


def build_ppt():
    prs = Presentation(PPT)
    slides = list(prs.slides)
    # 13쪽: '음식·음료' 행(이름표·0 / 12·0%·설명) 도형 삭제
    s13 = slides[12]
    kill = [sh for sh in s13.shapes if sh.has_text_frame and sh.text_frame.text.strip() in
            ("음식·음료", "0 / 12", "0%", "진열 상품·포장 그림 오인 → 지표에서 제외")]
    assert len(kill) == 4, [sh.text_frame.text for sh in kill]
    for sh in kill:
        sh._element.getparent().remove(sh._element)
    # 18쪽 한계 문구
    m12.set_par(slides[17], "좌석 정밀도 78%(일본 2/6) · 음식은 0%라 제외 · 흐린 장면 일부 포함",
                "좌석 정밀도 78%(일본 2/6) · 흐린 장면 일부 포함")
    # 16쪽 프로토타입 캡처 교체
    for sh in slides[15].shapes:
        if sh.shape_type == 13:
            part = sh.part.related_part(sh._element.blipFill.blip.rEmbed)
            f = MATCH["pptx"].get(part.partname.lstrip("/"))
            if f:
                part._blob = open(os.path.join(OUT, f), "rb").read()
    for sl, (_, _, t) in zip(slides, SCRIPT):
        tf = sl.notes_slide.notes_text_frame
        if tf is not None:
            tf.text = t
    prs.save(PPT)


SCRIPT = [list(r) for r in m12.SCRIPT]
SCRIPT[12][2] = SCRIPT[12][2].replace(" 반대로 음식은 열두 개 모두 진열 상품을 잘못 본 것이라 지표에서 뺐습니다.", "")
assert "음식" not in SCRIPT[12][2]
m12.SCRIPT[:] = [tuple(r) for r in SCRIPT]


def main():
    for f in (HWP_DST, PPT, m12.DOCX):
        m12.backup(f)
    build_hwp()
    build_ppt()
    m12.build_docx()
    print("[저장]", HWP_DST, PPT, m12.DOCX, sep="\n  ")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
